from io import BytesIO
from pathlib import PurePath

import pdfplumber
from docx import Document
from fastapi import HTTPException, UploadFile

from .config import settings

ALLOWED_EXTENSIONS = {".pdf", ".docx"}


async def read_cv(upload: UploadFile) -> tuple[str, bytes]:
    filename = PurePath(upload.filename or "").name
    extension = PurePath(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Upload a PDF or DOCX file")

    content = await upload.read()
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded CV is empty")
    if len(content) > settings.max_cv_bytes:
        raise HTTPException(status_code=413, detail="CV files must be smaller than 5 MB")

    try:
        if extension == ".pdf":
            text = _extract_pdf(content)
        else:
            text = _extract_docx(content)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="The CV could not be read") from exc

    text = text.strip()
    if not text:
        raise HTTPException(status_code=400, detail="The CV contains no readable text")
    return text[: settings.max_cv_text_chars], content


def _extract_pdf(content: bytes) -> str:
    with pdfplumber.open(BytesIO(content)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _extract_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)
